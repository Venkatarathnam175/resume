import streamlit as st
from PyPDF2 import PdfReader
from docx import Document
from fpdf import FPDF
from io import BytesIO

st.set_page_config(page_title="Free Resume Builder", layout="wide")


# ------------------------ EXTRACT RESUME TEXT ------------------------ #

def extract_text_from_pdf(file):
    reader = PdfReader(file)
    text = []
    for page in reader.pages:
        text.append(page.extract_text() or "")
    return "\n".join(text)


def extract_text_from_docx(file):
    doc = Document(file)
    return "\n".join([p.text for p in doc.paragraphs])


def parse_uploaded_resume(uploaded_file):
    if uploaded_file is None:
        return ""
    name = uploaded_file.name.lower()
    if name.endswith(".pdf"):
        return extract_text_from_pdf(uploaded_file)
    elif name.endswith(".docx"):
        return extract_text_from_docx(uploaded_file)
    return ""


# ------------------------ PDF GENERATOR (FPDF) ------------------------ #

def generate_pdf(data):
    pdf = FPDF(format='A4')
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    # HEADER
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, data['name'], ln=True)

    pdf.set_font("Arial", "", 11)
    pdf.multi_cell(0, 6,
                   f"{data['email']} | {data['phone']} | "
                   f"{data['location']} | {data['linkedin']} "
                   f"{'| ' + data['github'] if data['github'] else ''}"
                   )

    pdf.ln(2)

    # SUMMARY
    pdf.set_font("Arial", "B", 13)
    pdf.cell(0, 8, "SUMMARY", ln=True)
    pdf.set_font("Arial", "", 11)
    pdf.multi_cell(0, 6, data['summary'])
    pdf.ln(3)

    # EDUCATION
    pdf.set_font("Arial", "B", 13)
    pdf.cell(0, 8, "EDUCATION", ln=True)

    pdf.set_font("Arial", "", 11)
    for edu in data['education']:
        pdf.multi_cell(0, 6, f"{edu['degree']} - {edu['school']}")
        pdf.set_font("Arial", "I", 10)
        pdf.multi_cell(0, 5, f"{edu['start_end']} | {edu['location']}")
        pdf.set_font("Arial", "", 11)
        pdf.multi_cell(0, 6, edu['details'])
        pdf.ln(2)

    # EXPERIENCE
    pdf.set_font("Arial", "B", 13)
    pdf.cell(0, 8, "EXPERIENCE", ln=True)

    pdf.set_font("Arial", "", 11)
    for exp in data['experience']:
        pdf.multi_cell(0, 6, f"{exp['role']} - {exp['company']}")
        pdf.set_font("Arial", "I", 10)
        pdf.multi_cell(0, 5, f"{exp['start_end']} | {exp['location']}")
        pdf.set_font("Arial", "", 11)
        for b in exp['bullets'].split("\n"):
            if b.strip():
                pdf.multi_cell(0, 6, "• " + b.strip())
        pdf.ln(2)

    # PROJECTS
    pdf.set_font("Arial", "B", 13)
    pdf.cell(0, 8, "PROJECTS", ln=True)

    pdf.set_font("Arial", "", 11)
    for proj in data['projects']:
        pdf.multi_cell(0, 6, f"{proj['name']} ({proj['tech']})")
        for b in proj['bullets'].split("\n"):
            if b.strip():
                pdf.multi_cell(0, 6, "• " + b.strip())
        pdf.ln(2)

    # SKILLS
    pdf.set_font("Arial", "B", 13)
    pdf.cell(0, 8, "SKILLS", ln=True)

    pdf.set_font("Arial", "", 11)
    pdf.multi_cell(0, 6, f"Technical: {data['skills_technical']}")
    pdf.multi_cell(0, 6, f"Tools: {data['skills_tools']}")
    pdf.multi_cell(0, 6, f"Soft Skills: {data['skills_soft']}")

    return pdf.output(dest="S").encode("latin-1")


# ------------------------ MAIN UI ------------------------ #

st.title("🎓 Free Student Resume Builder (No Paywall)")

st.write("Fill the form → Preview → Download PDF. Upload old resume if needed.")

st.markdown("---")


# UPLOAD OLD RESUME
st.subheader("1️⃣ Upload Old Resume (optional)")
uploaded = st.file_uploader("Upload PDF or DOCX", type=["pdf", "docx"])
extracted = ""

if uploaded:
    extracted = parse_uploaded_resume(uploaded)
    st.success("Resume extracted successfully.")
    st.text_area("Extracted text (copy any part you want):", extracted, height=150)

st.markdown("---")


# ------------------------ INPUT FORM ------------------------ #

st.subheader("2️⃣ Enter Resume Details")

col1, col2 = st.columns(2)
with col1:
    name = st.text_input("Full Name")
    location = st.text_input("Location")
    email = st.text_input("Email")
with col2:
    phone = st.text_input("Phone")
    linkedin = st.text_input("LinkedIn URL")
    github = st.text_input("GitHub / Portfolio (optional)")

headline = st.text_input("Headline (e.g., B.Tech Student | Aspiring Analyst)")
summary = st.text_area("Summary (exact text):", height=100)

# EDUCATION
st.markdown("### Education")
edu_count = st.number_input("Number of education entries:", 1, 10, 1)
education = []

for i in range(edu_count):
    st.markdown(f"**Education #{i+1}**")
    c1, c2 = st.columns(2)
    with c1:
        degree = st.text_input(f"Degree #{i+1}")
        school = st.text_input(f"Institution #{i+1}")
    with c2:
        dates = st.text_input(f"Duration #{i+1}")
        loc = st.text_input(f"Location #{i+1}")
    details = st.text_area(f"Details #{i+1}", height=80)
    education.append({
        "degree": degree,
        "school": school,
        "start_end": dates,
        "location": loc,
        "details": details
    })

# EXPERIENCE
st.markdown("### Experience")
exp_count = st.number_input("Number of experiences:", 0, 15, 1)
experience = []
for i in range(exp_count):
    st.markdown(f"**Experience #{i+1}**")
    c1, c2 = st.columns(2)
    with c1:
        role = st.text_input(f"Role #{i+1}")
        company = st.text_input(f"Company #{i+1}")
    with c2:
        dates = st.text_input(f"Duration #{i+1}")
        loc = st.text_input(f"Location #{i+1}")
    bullets = st.text_area(f"Bullet points #{i+1}", height=100)
    experience.append({
        "role": role,
        "company": company,
        "start_end": dates,
        "location": loc,
        "bullets": bullets
    })

# PROJECTS
st.markdown("### Projects")
proj_count = st.number_input("Number of projects:", 0, 15, 1)
projects = []
for i in range(proj_count):
    st.markdown(f"**Project #{i+1}**")
    c1, c2 = st.columns(2)
    with c1:
        namep = st.text_input(f"Project Name #{i+1}")
    with c2:
        tech = st.text_input(f"Tech Used #{i+1}")
    bullets = st.text_area(f"Project bullets #{i+1}", height=100)
    projects.append({
        "name": namep,
        "tech": tech,
        "bullets": bullets
    })

# SKILLS
st.markdown("### Skills")
skills_technical = st.text_area("Technical Skills", height=60)
skills_tools = st.text_area("Tools / Software", height=60)
skills_soft = st.text_area("Soft Skills", height=60)

# COLLECT ALL DATA
resume = {
    "name": name,
    "headline": headline,
    "location": location,
    "email": email,
    "phone": phone,
    "linkedin": linkedin,
    "github": github,
    "summary": summary,
    "education": education,
    "experience": experience,
    "projects": projects,
    "skills_technical": skills_technical,
    "skills_tools": skills_tools,
    "skills_soft": skills_soft
}

st.markdown("---")

# PREVIEW IN HTML
if st.button("Generate HTML Preview"):
    st.session_state["preview_html"] = f"""
    <h1>{name}</h1>
    <p>{email} | {phone} | {location}</p>
    <hr>
    <h3>Summary</h3>
    <p>{summary}</p>
    """
    st.success("Preview created below!")

if "preview_html" in st.session_state:
    st.markdown("### Resume Preview")
    st.components.v1.html(st.session_state["preview_html"], height=500, scrolling=True)

# DOWNLOAD PDF
if st.button("Download PDF"):
    pdf_bytes = generate_pdf(resume)
    st.download_button("Download Resume PDF", pdf_bytes, "resume.pdf", "application/pdf")
